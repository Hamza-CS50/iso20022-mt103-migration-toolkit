"""Build the Excel dashboard + business case from pipeline output.  Run: python scripts/build_dashboard.py"""
import csv, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
from mt2iso.validator import RULES

OUT = "output"
rd = lambda n: list(csv.DictReader(open(os.path.join(OUT, n))))
results, findings, lineage, evalr, address = rd("results.csv"), rd("findings.csv"), rd("lineage.csv"), rd("evaluation.csv"), rd("address.csv")

F = "Arial"
base = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", start_color="1F3A5F")
sub_fill = PatternFill("solid", start_color="DCE6F1")
yellow = PatternFill("solid", start_color="FFFF00")
blue = Font(name=F, size=10, color="0000FF")
green = Font(name=F, size=10, color="008000")
title = Font(name=F, size=14, bold=True, color="1F3A5F")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0;($#,##0);-'
PCT = '0.0%'

wb = Workbook()


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, Alignment(horizontal="center", vertical="center", wrap_text=True), box


def put(ws, ref, v, font=base, fmt=None, fill=None, bd=True, wrap=False, align=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if bd: c.border = box
    if wrap or align: c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


def data_sheet(name, cols, rows, numeric=(), w=None):
    ws = wb.create_sheet(name)
    header(ws, 1, cols)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = base
    ws.freeze_panes = "A2"
    if w: widths(ws, w)
    ws.auto_filter.ref = f"A1:{L(len(cols))}{len(rows)+1}"
    return ws


# ============================================================ data sheets
wb.remove(wb.active)
ws_read = wb.create_sheet("README")
ws_sum = wb.create_sheet("Summary")

N = len(results)
res_cols = ["Msg ID", "Scenario", "Corridor", "Ccy", "Amount", "Injected defect", "Status", "Quality score", "Errors", "Warnings",
            "Inferences", "Data-loss flags", "Enrichments", "ISO leaf elements", "DIRECT", "CODE", "PARSED", "INFERRED", "FREETEXT", "Injected?"]
rows = []
for r in results:
    rows.append([r["msg_id"], r["scenario"], r["corridor"], r["currency"], float(r["amount"]), r["injected_defect"], r["status"], int(r["score"]),
                 int(r["errors"]), int(r["warnings"]), int(r["inference"]), int(r["data_loss"]), int(r["enrichment"]), int(r["leaf_elements"]),
                 int(r["direct"]), int(r["code"]), int(r["parsed"]), int(r["inferred"]), int(r["freetext"])])
ws_res = data_sheet("Results", res_cols, rows, w=[9, 9, 9, 6, 12, 26, 20, 8, 8, 9, 10, 10, 11, 10, 9, 9, 9, 10, 10, 10])
for r in range(2, N + 2):
    ws_res.cell(row=r, column=20, value=f'=IF(F{r}="","No","Yes")').font = base
    ws_res.cell(row=r, column=5).number_format = "#,##0.00"
for txt, col in (("PASS_CLEAN", "C6EFCE"), ("PASS_WITH_WARNINGS", "FFEB9C"), ("FAILED_VALIDATION", "F8CBAD"), ("REJECTED", "E6B8B7")):
    ws_res.conditional_formatting.add(f"G2:G{N+1}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", start_color=col, end_color=col)))
RES = lambda col: f"Results!${col}$2:${col}${N+1}"

NF = len(findings)
ws_f = data_sheet("Findings", ["Msg ID", "Rule ID", "Category", "Severity", "Element / field", "Message"],
                  [[f["msg_id"], f["rule_id"], f["category"], f["severity"], f["path"], f["message"]] for f in findings], w=[9, 26, 12, 10, 44, 100])
FND = lambda col: f"Findings!${col}$2:${col}${NF+1}"

NL = len(lineage)
data_sheet("Lineage Data", ["Msg ID", "ISO 20022 element path", "Group", "Source in MT103", "Method"],
           [[l["msg_id"], l["path"], l["group"], l["source"], l["method"]] for l in lineage], w=[9, 50, 16, 46, 11])
LIN = lambda col: f"'Lineage Data'!${col}$2:${col}${NL+1}"

NA = len(address)
ws_ad = data_sheet("Address Data", ["Msg ID", "Party role", "Country", "Town+Country present (1/0)", "Street structured (1/0)", "Has AdrLine (1/0)"],
                   [[a["msg_id"], a["role"], a["country"], int(a["town_ctry"] == "True"), int(a["street"] == "True"), int(a["adrline"] == "True")] for a in address], w=[9, 10, 9, 16, 16, 14])
ADR = lambda col: f"'Address Data'!${col}$2:${col}${NA+1}"

# ---------------- Detection Eval
ws_ev = wb.create_sheet("Detection Eval")
header(ws_ev, 1, ["Msg ID", "Injected defect", "Expected rule", "Detected (1/0)"])
for i, e in enumerate(evalr, 2):
    ws_ev.cell(row=i, column=1, value=e["msg_id"]).font = base
    ws_ev.cell(row=i, column=2, value=e["defect"]).font = base
    ws_ev.cell(row=i, column=3, value=e["expected_rule"]).font = base
    ws_ev.cell(row=i, column=4, value=f'=IF(COUNTIFS({FND("A")},A{i},{FND("B")},C{i})>0,1,0)').font = base
NE = len(evalr)
defects = sorted({e["defect"] for e in evalr})
header(ws_ev, 1, ["Defect type", "Injected", "Detected", "Recall"], col=6)
for i, d in enumerate(defects, 2):
    put(ws_ev, f"F{i}", d)
    put(ws_ev, f"G{i}", f"=COUNTIF($B$2:$B${NE+1},F{i})")
    put(ws_ev, f"H{i}", f"=SUMIF($B$2:$B${NE+1},F{i},$D$2:$D${NE+1})")
    put(ws_ev, f"I{i}", f"=IF(G{i}=0,0,H{i}/G{i})", fmt=PCT)
tr = len(defects) + 2
put(ws_ev, f"F{tr}", "TOTAL", bold, fill=sub_fill); put(ws_ev, f"G{tr}", f"=SUM(G2:G{tr-1})", bold, fill=sub_fill)
put(ws_ev, f"H{tr}", f"=SUM(H2:H{tr-1})", bold, fill=sub_fill); put(ws_ev, f"I{tr}", f"=IF(G{tr}=0,0,H{tr}/G{tr})", bold, PCT, sub_fill)
widths(ws_ev, [9, 28, 26, 12, 3, 30, 10, 10, 10])
ws_ev.freeze_panes = "A2"
ch = BarChart(); ch.type = "bar"; ch.title = "Detection recall by injected defect"; ch.style = 10
ch.add_data(Reference(ws_ev, min_col=9, min_row=1, max_row=tr - 1), titles_from_data=True)
ch.set_categories(Reference(ws_ev, min_col=6, min_row=2, max_row=tr - 1))
ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.number_format = 0, 1, "0%"
ch.legend = None; ch.height, ch.width = 10, 18
ws_ev.add_chart(ch, f"F{tr+3}")

# ---------------- Rule Catalog
ws_rc = wb.create_sheet("Rule Catalog")
header(ws_rc, 1, ["Rule ID", "Category", "Severity", "What it checks", "Times fired (this run)"])
for i, (rid, (cat, sev, desc)) in enumerate(RULES.items(), 2):
    for c, v in enumerate([rid, cat, sev, desc], 1):
        cell = ws_rc.cell(row=i, column=c, value=v); cell.font = base; cell.border = box
    put(ws_rc, f"E{i}", f"=COUNTIF({FND('B')},A{i})")
widths(ws_rc, [26, 12, 10, 90, 16]); ws_rc.freeze_panes = "A2"
NR = len(RULES)

# ---------------- Mapping matrix
MM = [
 (":20:", "Sender's reference", "GrpHdr/MsgId; PmtId/InstrId; PmtId/TxId", "DIRECT", "16 chars", "35 chars per ID", "NEUTRAL", "Same value reused for three IDs"),
 (":23B:", "Bank operation code", "PmtTpInf/SvcLvl/Prtry (SPRI, SSTD, SPAY)", "CODE", "5 codes", "Service level / category purpose", "NEUTRAL", "CRED is the default and produces no element. Mapping decision of this toolkit; CBPR+ guidelines may differ"),
 (":32A:", "Value date, currency, amount", "IntrBkSttlmDt; IntrBkSttlmAmt (+Ccy)", "DIRECT", "Decimal comma; YYMMDD", "xs:date, decimal point, 18 digits", "NEUTRAL", "Format conversion only; decimals checked per ISO 4217"),
 (":33B:", "Instructed amount", "InstdAmt (+Ccy)", "DIRECT", "Optional", "Native element", "NEUTRAL", ""),
 (":36:", "Exchange rate", "XchgRate", "DIRECT", "Decimal comma", "Decimal point, up to 10 decimals", "NEUTRAL", "Rule: 33B currency <> 32A currency should come with a rate"),
 (":50A:", "Ordering customer (BIC)", "Dbtr/Id/OrgId/AnyBIC", "DIRECT", "BIC only, no name", "Name + address + IDs", "LOSS", "No name available, flagged by BR_NAME_REQ"),
 (":50K:", "Ordering customer (free text)", "Dbtr/Nm; Dbtr/PstlAdr/*", "PARSED", "4x35 chars, name+address mixed", "Nm 140; structured address; 16+ elements", "GUESS", "Address structured by heuristics; leftovers go to AdrLine (hybrid)"),
 (":50F:", "Ordering customer (structured)", "Dbtr/Nm; Dbtr/PstlAdr/TwnNm,Ctry,AdrLine", "DIRECT", "Numbered lines 1/-8/", "Native structure", "GAIN", "Lines 4/-8/ not carried -> DL_UNMAPPED_50F_LINES"),
 (":50a: account", "Ordering account", "DbtrAcct/Id/IBAN or Othr/Id", "DIRECT", "Account in first line", "IBAN vs other account typed", "GAIN", "IBAN checksum and length validated"),
 (":52A/D:", "Ordering institution", "DbtrAgt/FinInstnId/BICFI or Nm+PstlAdr", "DIRECT", "Optional", "Mandatory in pacs.008", "GUESS", "If absent the sending bank BIC is assumed (INF_DBTRAGT_INFERRED)"),
 (":53A:", "Sender's correspondent", "SttlmInf/InstgRmbrsmntAgt; SttlmMtd=INGA", "DIRECT", "Optional", "Settlement instruction block", "NEUTRAL", "Settlement-method choice is simplified"),
 (":54A:", "Receiver's correspondent", "SttlmInf/InstdRmbrsmntAgt", "DIRECT", "Optional", "Settlement instruction block", "NEUTRAL", ""),
 (":56A/D:", "Intermediary institution", "IntrmyAgt1", "DIRECT", "Optional", "Up to 3 intermediaries", "NEUTRAL", ""),
 (":57A/D:", "Account with institution", "CdtrAgt/FinInstnId", "DIRECT", "Optional", "Mandatory in pacs.008", "GUESS", "If absent the receiving bank BIC is assumed (INF_CDTRAGT_INFERRED)"),
 (":59:", "Beneficiary (free text)", "Cdtr/Nm; Cdtr/PstlAdr/*; CdtrAcct", "PARSED", "4x35 chars; name+address mixed", "Structured address", "GUESS", "Same heuristics as 50K"),
 (":59F:", "Beneficiary (structured)", "Cdtr/Nm; Cdtr/PstlAdr/*", "DIRECT", "Numbered lines", "Native structure", "GAIN", ""),
 (":70:", "Remittance information", "RmtInf/Ustrd; RmtInf/Strd/RfrdDocInf; CdtrRefInf", "CODE / FREETEXT", "4x35 = 140 chars", "Ustrd 140 + structured blocks", "GAIN", "/INV/ -> CINV document number; /RFB/ -> creditor reference; /ROC/ -> EndToEndId"),
 (":70: /ROC/", "Ordering customer's reference", "PmtId/EndToEndId", "CODE", "Buried in free text", "Dedicated element", "GAIN", "Absent -> NOTPROVIDED (DL / BR_E2E_NOTPROVIDED)"),
 (":71A:", "Details of charges", "ChrgBr (OUR=DEBT, BEN=CRED, SHA=SHAR)", "CODE", "3 codes", "4 codes (adds SLEV)", "NEUTRAL", "Charge-combination rules checked at MT level"),
 (":71F:/:71G:", "Sender's / receiver's charges", "ChrgsInf/Amt + Agt", "DIRECT / INFERRED", "No bank named", "Agent mandatory", "GUESS", "Charging agent assumed (INF_CHRG_AGENT_INFERRED)"),
 (":72:", "Sender to receiver info", "InstrForNxtAgt/InstrInf", "FREETEXT", "6x35 with code words", "Coded instruction options", "NEUTRAL", "Stays unstructured here"),
 (":77B:", "Regulatory reporting", "RgltryRptg/Dtls/Inf", "FREETEXT", "3x35", "Structured regulatory details", "NEUTRAL", "Could be structured further (authority, country, code)"),
 ("{3:{121:}}", "UETR (gpi tracker)", "PmtId/UETR", "DIRECT / INFERRED", "Optional", "UUIDv4", "GUESS", "If absent a UETR is minted (INF_UETR_GENERATED)"),
 ("{1:}/{2:}", "Sender / receiver LT", "InstgAgt, InstdAgt BICFI", "DIRECT", "LT address", "Agent elements", "NEUTRAL", "BIC8 + branch extracted from the 12-char LT address"),
]
ws_mm = wb.create_sheet("Mapping Matrix")
header(ws_mm, 1, ["MT103 field", "Meaning", "ISO 20022 target (pacs.008.001.08)", "Method", "MT constraint", "ISO 20022 capability", "Effect", "Notes"])
for i, row in enumerate(MM, 2):
    for c, v in enumerate(row, 1):
        cell = ws_mm.cell(row=i, column=c, value=v); cell.font = base; cell.border = box; cell.alignment = Alignment(wrap_text=True, vertical="top")
for eff, col in (("GAIN", "C6EFCE"), ("GUESS", "FFEB9C"), ("LOSS", "F8CBAD")):
    ws_mm.conditional_formatting.add(f"G2:G{len(MM)+1}", CellIsRule(operator="equal", formula=[f'"{eff}"'], fill=PatternFill("solid", start_color=col, end_color=col)))
widths(ws_mm, [14, 28, 42, 16, 26, 30, 10, 60]); ws_mm.freeze_panes = "A2"

# ---------------- Lineage summary
ws_ln = wb.create_sheet("Lineage")
put(ws_ln, "A1", "Where did each ISO 20022 value come from?", title, bd=False)
put(ws_ln, "A2", "Counts of populated leaf elements across all converted messages (live COUNTIFS over 'Lineage Data').", base, bd=False)
methods = ["DIRECT", "CODE", "PARSED", "INFERRED", "FREETEXT"]
groups = sorted({l["group"] for l in lineage})
header(ws_ln, 4, ["Element group"] + methods + ["Total", "Deterministic %", "Guessed % (PARSED+INFERRED)"])
for i, g in enumerate(groups, 5):
    put(ws_ln, f"A{i}", g)
    for j, m in enumerate(methods):
        put(ws_ln, f"{L(2+j)}{i}", f'=COUNTIFS({LIN("C")},$A{i},{LIN("E")},{L(2+j)}$4)')
    put(ws_ln, f"G{i}", f"=SUM(B{i}:F{i})")
    put(ws_ln, f"H{i}", f"=IF(G{i}=0,0,(B{i}+C{i})/G{i})", fmt=PCT)
    put(ws_ln, f"I{i}", f"=IF(G{i}=0,0,(D{i}+E{i})/G{i})", fmt=PCT)
lt = 5 + len(groups)
put(ws_ln, f"A{lt}", "TOTAL", bold, fill=sub_fill)
for j in range(2, 8):
    put(ws_ln, f"{L(j)}{lt}", f"=SUM({L(j)}5:{L(j)}{lt-1})", bold, fill=sub_fill)
put(ws_ln, f"H{lt}", f"=IF(G{lt}=0,0,(B{lt}+C{lt})/G{lt})", bold, PCT, sub_fill)
put(ws_ln, f"I{lt}", f"=IF(G{lt}=0,0,(D{lt}+E{lt})/G{lt})", bold, PCT, sub_fill)
widths(ws_ln, [20, 10, 10, 10, 10, 10, 10, 16, 24])
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Value provenance by element group"
ch.add_data(Reference(ws_ln, min_col=2, max_col=6, min_row=4, max_row=lt - 1), titles_from_data=True)
ch.set_categories(Reference(ws_ln, min_col=1, min_row=5, max_row=lt - 1)); ch.height, ch.width = 9, 22
ws_ln.add_chart(ch, f"A{lt+3}")

# ============================================================ Summary
ws = ws_sum
put(ws, "A1", "ISO 20022 Migration Toolkit - Results Dashboard", title, bd=False)
put(ws, "A2", "Every figure below is a live formula over the Results / Findings / Lineage / Address sheets (output of the Python pipeline on the synthetic stress-test dataset).", base, bd=False)
put(ws, "A4", "Headline KPIs", bold, bd=False)
kp = [
 ("Messages processed", f"=COUNTA({RES('A')})", "0"),
 ("Converted & clean (no findings above INFO)", f'=COUNTIF({RES("G")},"PASS_CLEAN")', "0"),
 ("Converted with warnings", f'=COUNTIF({RES("G")},"PASS_WITH_WARNINGS")', "0"),
 ("Converted but failed validation (needs repair)", f'=COUNTIF({RES("G")},"FAILED_VALIDATION")', "0"),
 ("Rejected at parse (no ISO message produced)", f'=COUNTIF({RES("G")},"REJECTED")', "0"),
 ("Straight-through rate (clean + warnings)", "=(B6+B7)/B5", PCT),
 ("Repair rate (failed + rejected)", "=(B8+B9)/B5", PCT),
 ("Average quality score (0-100)", f"=AVERAGE({RES('H')})", "0.0"),
 ("ISO leaf elements: deterministic (DIRECT+CODE)", f"=(SUM({RES('O')})+SUM({RES('P')}))/SUM({RES('N')})", PCT),
 ("ISO leaf elements: heuristic (PARSED)", f"=SUM({RES('Q')})/SUM({RES('N')})", PCT),
 ("ISO leaf elements: inferred (not in MT103)", f"=SUM({RES('R')})/SUM({RES('N')})", PCT),
 ("ISO leaf elements: free text carried over", f"=SUM({RES('S')})/SUM({RES('N')})", PCT),
 ("Injected defects", f'=COUNTIF({RES("T")},"Yes")', "0"),
 ("Injected defects detected by expected rule", f"='Detection Eval'!H{tr}", "0"),
 ("Detection recall", "=IF(B17=0,0,B18/B17)", PCT),
 ("False alarms: messages with NO injected defect but an ERROR", f'=COUNTIFS({RES("T")},"No",{RES("I")},">0")', "0"),
 ("Parties with town+country in output (CBPR+ Nov-2026 ready)", f"=SUM({ADR('D')})/COUNTA({ADR('A')})", PCT),
 ("Parties with street/building structured", f"=SUM({ADR('E')})/COUNTA({ADR('A')})", PCT),
]
for i, (lab, f, fmt) in enumerate(kp, 5):
    put(ws, f"A{i}", lab); put(ws, f"B{i}", f, bold, fmt)
r0 = 5 + len(kp) + 1   # status table
put(ws, f"A{r0}", "Status distribution", bold, bd=False)
header(ws, r0 + 1, ["Status", "Messages", "Share"])
for i, st in enumerate(["PASS_CLEAN", "PASS_WITH_WARNINGS", "FAILED_VALIDATION", "REJECTED"], r0 + 2):
    put(ws, f"A{i}", st); put(ws, f"B{i}", f'=COUNTIF({RES("G")},A{i})'); put(ws, f"C{i}", f"=B{i}/$B$5", fmt=PCT)
sr = r0 + 2
c1 = BarChart(); c1.type = "col"; c1.title = "Migration outcome per message"; c1.style = 10; c1.legend = None
c1.add_data(Reference(ws, min_col=2, min_row=r0 + 1, max_row=r0 + 5), titles_from_data=True)
c1.set_categories(Reference(ws, min_col=1, min_row=sr, max_row=sr + 3)); c1.height, c1.width = 7.5, 13
ws.add_chart(c1, "E4")

r1 = r0 + 8   # category table
put(ws, f"A{r1}", "Findings by category", bold, bd=False)
header(ws, r1 + 1, ["Category", "Findings", "Share"])
cats = ["SYNTAX", "SCHEMA", "BUSINESS", "ADDRESS", "DATA_LOSS", "INFERENCE", "ENRICHMENT"]
for i, c in enumerate(cats, r1 + 2):
    put(ws, f"A{i}", c); put(ws, f"B{i}", f'=COUNTIF({FND("C")},A{i})'); put(ws, f"C{i}", f"=B{i}/COUNTA({FND('A')})", fmt=PCT)
c2 = BarChart(); c2.type = "bar"; c2.title = "Findings by category"; c2.style = 10; c2.legend = None
c2.add_data(Reference(ws, min_col=2, min_row=r1 + 1, max_row=r1 + 1 + len(cats)), titles_from_data=True)
c2.set_categories(Reference(ws, min_col=1, min_row=r1 + 2, max_row=r1 + 1 + len(cats))); c2.height, c2.width = 7.5, 13
ws.add_chart(c2, f"E{r0+3}")

r2 = r1 + len(cats) + 4   # severity + scenario
put(ws, f"A{r2}", "Outcome by business scenario", bold, bd=False)
header(ws, r2 + 1, ["Scenario", "Messages", "Clean", "Warnings", "Failed", "Rejected", "Repair rate"])
SCN = [("S1", "S1 Inbound worker remittance"), ("S2", "S2 Inbound trade"), ("S3", "S3 Outbound trade"), ("S4", "S4 Services")]
for i, (code, name) in enumerate(SCN, r2 + 2):
    put(ws, f"A{i}", code)
    put(ws, f"B{i}", f'=COUNTIF({RES("B")},A{i})')
    for j, st in enumerate(["PASS_CLEAN", "PASS_WITH_WARNINGS", "FAILED_VALIDATION", "REJECTED"]):
        put(ws, f"{L(3+j)}{i}", f'=COUNTIFS({RES("B")},$A{i},{RES("G")},"{st}")')
    put(ws, f"G{i}", f"=IF(B{i}=0,0,(E{i}+F{i})/B{i})", fmt=PCT)
put(ws, f"I{r2+1}", "Scenario key", bold, bd=False)
for i, (code, name) in enumerate(SCN, r2 + 2):
    put(ws, f"I{i}", name, bd=False)

r3 = r2 + 8   # address table
put(ws, f"A{r3}", "Address structuring by party country (heuristic parser)", bold, bd=False)
header(ws, r3 + 1, ["Country", "Parties", "Town+Country found", "Street structured", "Town+Country %", "Street %"])
countries = sorted({a["country"] for a in address})
for i, c in enumerate(countries, r3 + 2):
    put(ws, f"A{i}", c)
    put(ws, f"B{i}", f'=COUNTIF({ADR("C")},A{i})')
    put(ws, f"C{i}", f'=SUMIF({ADR("C")},A{i},{ADR("D")})')
    put(ws, f"D{i}", f'=SUMIF({ADR("C")},A{i},{ADR("E")})')
    put(ws, f"E{i}", f"=IF(B{i}=0,0,C{i}/B{i})", fmt=PCT); put(ws, f"F{i}", f"=IF(B{i}=0,0,D{i}/B{i})", fmt=PCT)
widths(ws, [58, 14, 20, 18, 16, 12, 12, 3, 34])

# ============================================================ Business case
bc = wb.create_sheet("Business Case")
put(bc, "A1", "Business case - migrating a mid-size bank's cross-border flows to ISO 20022 data quality", title, bd=False)
put(bc, "A2", "ALL inputs are ILLUSTRATIVE assumptions (blue). Replace them with the bank's own volumes, repair logs and screening statistics. They are not derived from the toolkit's synthetic data.", Font(name=F, size=10, italic=True, color="C00000"), bd=False)
put(bc, "A3", "Scenario selector (1 = Conservative, 2 = Base, 3 = Optimistic)  >>", bold, bd=False)
put(bc, "B3", 2, blue, fill=yellow)
header(bc, 4, ["Input", "Conservative", "Base", "Optimistic", "ACTIVE", "Unit / how to read"])
INP = [
 ("Cross-border payments per year", 400000, 600000, 800000, "messages", "#,##0"),
 ("Manual repair rate today", 0.025, 0.04, 0.055, "share of messages touched by an operator", PCT),
 ("Relative cut in repairs after migration", 0.15, 0.30, 0.45, "structured data reduces free-text fixing", PCT),
 ("Cost per manual repair", 8, 12, 16, "USD per repaired message (staff time)", USD),
 ("Sanctions-screening alert rate", 0.02, 0.03, 0.04, "share of messages generating an alert", PCT),
 ("Relative cut in false-positive alerts", 0.05, 0.12, 0.20, "better name/address fields reduce fuzzy matches", PCT),
 ("Cost per alert review", 6, 10, 14, "USD per alert (analyst time)", USD),
 ("One-off migration cost", 900000, 650000, 500000, "USD: mapping, integration, testing, training (year 0)", USD),
 ("Annual run cost", 150000, 110000, 80000, "USD per year: licences, support, monitoring", USD),
 ("Discount rate", 0.12, 0.10, 0.08, "cost of capital", PCT),
 ("Volume growth", 0.02, 0.05, 0.08, "per year", PCT),
 ("Payments with unstructured address today", 0.40, 0.58, 0.60, "share. Base echoes Swift's published network-wide July 2026 level (about 58-59%); a bank's own share will differ", PCT),
 ("Share of those NAK'd if not remediated", 0.02, 0.05, 0.10, "payments Swift would reject from 14 Nov 2026 and not fixed upstream", PCT),
 ("Cost per NAK'd payment", 20, 30, 45, "USD: investigation, resubmission, customer compensation", USD),
]
for i, (lab, a, b_, c, unit, fmt) in enumerate(INP, 5):
    put(bc, f"A{i}", lab)
    for col, v in zip("BCD", (a, b_, c)):
        put(bc, f"{col}{i}", v, blue, fmt)
    put(bc, f"E{i}", f"=CHOOSE($B$3,B{i},C{i},D{i})", bold, fmt)
    put(bc, f"F{i}", unit)
# inputs: E5 volume, E6 repair rate, E7 cut, E8 cost/repair, E9 alert rate, E10 FP cut, E11 cost/alert,
#         E12 one-off, E13 run, E14 discount, E15 growth, E16 unstructured share, E17 NAK share, E18 cost/NAK
put(bc, "A20", "Five-year cash-flow model (USD)", bold, bd=False)
header(bc, 21, ["Year", "0", "1", "2", "3", "4", "5"])
put(bc, "A22", "Year index", bold)
for j, col in enumerate("BCDEFG"):
    put(bc, f"{col}22", j, bold)
def line(r, lab, fn, fmt, b_=False):
    put(bc, f"A{r}", lab, bold if b_ else base)
    put(bc, f"B{r}", None)
    for col in "CDEFG":
        put(bc, f"{col}{r}", fn(col), bold if b_ else base, fmt)
line(23, "Payments (messages)", lambda c: f"=$E$5*(1+$E$15)^({c}$22-1)", "#,##0")
line(24, "Operational saving: fewer manual repairs", lambda c: f"={c}23*$E$6*$E$7*$E$8", USD)
line(25, "Operational saving: fewer false-positive alerts", lambda c: f"={c}23*$E$9*$E$10*$E$11", USD)
line(26, "Avoided cost: payments no longer NAK'd", lambda c: f"={c}23*$E$16*$E$17*$E$18", USD)
line(27, "Total benefit", lambda c: f"={c}24+{c}25+{c}26", USD, True)
line(28, "Run cost", lambda c: "=-$E$13", USD)
put(bc, "A29", "One-off migration cost"); put(bc, "B29", "=-$E$12", base, USD)
for col in "CDEFG": put(bc, f"{col}29", 0, base, USD)
put(bc, "A30", "Net cash flow", bold); put(bc, "B30", "=B29", bold, USD)
for col in "CDEFG": put(bc, f"{col}30", f"={col}27+{col}28+{col}29", bold, USD)
put(bc, "A31", "Discount factor")
for col in "BCDEFG": put(bc, f"{col}31", f"=1/(1+$E$14)^{col}22", base, "0.000")
put(bc, "A32", "Present value of net cash flow")
for col in "BCDEFG": put(bc, f"{col}32", f"={col}30*{col}31", base, USD)
put(bc, "A33", "Cumulative net cash flow (undiscounted)", bold); put(bc, "B33", "=B30", bold, USD)
prev = "B"
for col in "CDEFG":
    put(bc, f"{col}33", f"={prev}33+{col}30", bold, USD); prev = col
put(bc, "A34", "Payback flag (1 = cumulative >= 0)"); put(bc, "B34", 0)
for col in "CDEFG": put(bc, f"{col}34", f"=IF({col}33>=0,1,0)")
put(bc, "A36", "Key results", bold, bd=False)
KR = [("Net present value, all benefits (USD)", "=SUM(B32:G32)", USD),
      ("Net present value, operational savings ONLY (USD)", "=-$E$12+SUMPRODUCT((C24:G24+C25:G25+C28:G28)*C31:G31)", USD),
      ("IRR (all benefits)", "=IFERROR(IRR(B30:G30),\"n/a\")", PCT),
      ("Payback year (cumulative cash turns positive)", "=IFERROR(MATCH(1,C34:G34,0),\"beyond year 5\")", "0"),
      ("5-year ROI ((benefits - costs) / costs)", "=(SUM(C27:G27)-($E$12+5*$E$13))/($E$12+5*$E$13)", PCT),
      ("Share of 5-year benefits that is avoided NAK cost", "=SUM(C26:G26)/SUM(C27:G27)", PCT)]
for i, (lab, f, fmt) in enumerate(KR, 37):
    put(bc, f"A{i}", lab); put(bc, f"B{i}", f, bold, fmt)
put(bc, "A44", "Reference only: repair rate in the toolkit's synthetic stress test (46% of messages carry an injected defect; NOT a prediction for real flows)", Font(name=F, size=9, italic=True), bd=False)
put(bc, "A45", "Stress-test repair rate"); put(bc, "B45", "=Summary!B11", green, PCT)

put(bc, "A47", "Sensitivity A - NPV (USD): annual volume (rows) vs relative cut in repairs (columns); all other inputs = ACTIVE", bold, bd=False)
put(bc, "A48", "Volume \\ repair cut", bold, fill=sub_fill)
cuts = [0.10, 0.20, 0.30, 0.40, 0.50]
vols = [200000, 400000, 600000, 800000, 1000000]
for j, cval in enumerate(cuts):
    put(bc, f"{L(2+j)}48", cval, blue, PCT, sub_fill)
UNIT = "($E$6*{c}$48*$E$8+$E$9*$E$10*$E$11+$E$16*$E$17*$E$18)"
for i, v in enumerate(vols, 49):
    put(bc, f"A{i}", v, blue, "#,##0", sub_fill)
    for j in range(5):
        col = L(2 + j)
        put(bc, f"{col}{i}", f"=-$E$12+SUMPRODUCT(($A{i}*(1+$E$15)^($C$22:$G$22-1)*{UNIT.format(c=col)}-$E$13)/(1+$E$14)^$C$22:$G$22)", base, USD)
put(bc, "A55", "Sensitivity B - NPV (USD): share of unstructured-address payments NAK'd (rows) vs cost per NAK'd payment (columns); all other inputs = ACTIVE", bold, bd=False)
put(bc, "A56", "NAK share \\ cost per NAK", bold, fill=sub_fill)
naks = [0.0, 0.01, 0.02, 0.05, 0.10]
costs = [10, 20, 30, 45, 60]
for j, cval in enumerate(costs):
    put(bc, f"{L(2+j)}56", cval, blue, USD, sub_fill)
for i, v in enumerate(naks, 57):
    put(bc, f"A{i}", v, blue, PCT, sub_fill)
    for j in range(5):
        col = L(2 + j)
        put(bc, f"{col}{i}", f"=-$E$12+SUMPRODUCT(($E$5*(1+$E$15)^($C$22:$G$22-1)*($E$6*$E$7*$E$8+$E$9*$E$10*$E$11+$E$16*$A{i}*{col}$56)-$E$13)/(1+$E$14)^$C$22:$G$22)", base, USD)
for rng in ("B49:F53", "B57:F61"):
    bc.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")))
    bc.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")))
put(bc, "A63", "Colour code: blue = input, black = formula, green = link to another sheet, yellow = scenario selector. Reading the result: operational savings alone rarely repay a migration; the case rests on avoiding rejected payments after the Swift deadline, which is why that assumption is isolated and stress-tested in Sensitivity B.", Font(name=F, size=9, italic=True), bd=False)
widths(bc, [58, 16, 16, 16, 16, 60, 16]); bc.freeze_panes = "B5"
lc = LineChart(); lc.title = "Cumulative net cash flow (ACTIVE scenario, USD)"; lc.height, lc.width = 8, 16
lc.add_data(Reference(bc, min_col=1, max_col=7, min_row=33), from_rows=True, titles_from_data=True)
lc.set_categories(Reference(bc, min_col=2, max_col=7, min_row=22)); lc.legend = None
bc.add_chart(lc, "H3")

# ============================================================ README
rw = ws_read
put(rw, "A1", "ISO 20022 Legacy-to-MX Migration Toolkit - Excel companion", title, bd=False)
txt = [
 "What this is: the analysis workbook for the Python toolkit that converts SWIFT MT103 payments into ISO 20022 pacs.008.001.08, validates them, and measures what the migration gains, loses or guesses.",
 "Data: 100 fully synthetic MT103 messages (Pakistan-centred corridors). 46 carry exactly one injected defect so detection can be measured. All names, accounts and BICs are invented.",
 "",
 "Sheets",
 "  Summary - headline KPIs, status/category charts, scenario and address tables (all formulas).",
 "  Results - one row per message (toolkit output). Findings - every rule hit. Lineage Data / Address Data - raw rows behind the pivots.",
 "  Detection Eval - ground-truth check: was each injected defect caught by the rule that should catch it? Recall by defect type.",
 "  Rule Catalog - all rules with severity and how often each fired. Mapping Matrix - MT103 field -> ISO 20022 element, with gain / guess / loss.",
 "  Lineage - provenance of every ISO value (DIRECT, CODE, PARSED, INFERRED, FREETEXT). Business Case - 5-year NPV/IRR model with sensitivity.",
 "",
 "Colour code (financial-model convention): blue text = input you may change; black = formula; green = link to another sheet; yellow fill = scenario selector.",
 "To re-run the pipeline: python -m mt2iso generate && python -m mt2iso batch && python scripts/build_dashboard.py",
 "Important limits: the dataset is a stress test, not a sample of real traffic; the XSD is a subset authored from the published pacs.008.001.08 structure (validate against the official XSD before production);",
 "business-case inputs are illustrative placeholders and must be replaced with bank data.",
 "Source for the deadline referenced in rule BR_ADDR_TWN_CTRY: Swift, 'Removal of unstructured address' (unstructured addresses NAK'd under CBPR+ from 14 November 2026).",
]
for i, t in enumerate(txt, 3):
    put(rw, f"A{i}", t, bold if t == "Sheets" else base, bd=False, wrap=False)
rw.column_dimensions["A"].width = 200

order = ["README", "Summary", "Business Case", "Results", "Findings", "Detection Eval", "Rule Catalog", "Mapping Matrix", "Lineage", "Lineage Data", "Address Data"]
wb._sheets = [wb[n] for n in order]
wb.save(os.path.join(OUT, "ISO20022_Migration_Dashboard.xlsx"))
print("saved")
